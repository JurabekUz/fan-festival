"""Moderator eksportlari uchun bezatilgan Excel (.xlsx) fayllar."""
import io
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

TZ = timedelta(hours=5)  # Toshkent (UTC+5)
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
HEAD_FILL = PatternFill("solid", fgColor="1B2554")
HEAD_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
STRIPE = PatternFill("solid", fgColor="F3F4FA")
THIN = Side(style="thin", color="C9CDE4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def local_time(iso: str | None):
    """'2026-10-21T05:16:00Z' -> Toshkent vaqtidagi datetime (Excel sana sifatida)."""
    if not iso:
        return None
    return datetime.strptime(iso[:19], "%Y-%m-%dT%H:%M:%S") + TZ


def build(title: str, header: list[str], rows: list[list], time_cols=(), center_cols=()) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(header)
    for r in rows:
        ws.append([local_time(v) if i in time_cols else v for i, v in enumerate(r)])

    for c in ws[1]:
        c.font, c.fill, c.alignment, c.border = HEAD_FONT, HEAD_FILL, HEAD_ALIGN, BORDER
    ws.row_dimensions[1].height = 30
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BORDER
            if c.row % 2 == 0:
                c.fill = STRIPE
            if c.data_type == "f":  # foydalanuvchi matni hech qachon formula bo'lmasin
                c.data_type = "s"
            if c.column - 1 in time_cols:
                c.number_format = "dd.mm.yyyy hh:mm"
            if c.column - 1 in center_cols:
                c.alignment = Alignment(horizontal="center")

    for i, h in enumerate(header):
        longest = max([len(str(h))] + [16 if i in time_cols else len(str(r[i] if r[i] is not None else ""))
                                       for r in rows[:2000]])
        ws.column_dimensions[get_column_letter(i + 1)].width = min(max(longest + 3, 8), 45)
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
