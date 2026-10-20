from fastapi.testclient import TestClient

from conftest import register


def test_register_login_persist(client):
    c = register(client, 1, stamps=3)
    me = c.get("/api/me").json()
    assert me["stamps"] == [1, 2, 3] and me["passport_no"] == "F26-0001"
    # boshqa qurilmadan kirish — muhrlar joyida
    from app.main import app
    c2 = TestClient(app)
    assert c2.post("/api/login", json={"phone": "+998901000001", "pin": "1234"}).status_code == 200
    assert c2.get("/api/me").json()["stamps"] == [1, 2, 3]


def test_phone_unique_and_validation(client):
    register(client, 1, stamps=0)
    r = client.post("/api/register", json={"first_name": "A", "last_name": "B", "phone": "+998 90 100 0001", "pin": "1234", "visitor_type_id": 1})
    assert r.status_code == 409
    assert client.post("/api/register", json={"first_name": "A", "last_name": "B", "phone": "123", "pin": "1234"}).status_code == 422
    assert client.post("/api/register", json={"first_name": "A", "last_name": "B", "phone": "901112233", "pin": "12a4"}).status_code == 422


def test_wrong_code_and_rate_limit(client):
    c = register(client, 1, stamps=0)
    for _ in range(5):
        assert c.post("/api/checkin", json={"zone": 1, "code": "0000"}).status_code == 400
    assert c.post("/api/checkin", json={"zone": 1, "code": "1111"}).status_code == 429


def test_login_lockout(client):
    register(client, 1, stamps=0)
    from app.main import app
    c = TestClient(app)
    for _ in range(5):
        assert c.post("/api/login", json={"phone": "901000001", "pin": "9999"}).status_code == 401
    assert c.post("/api/login", json={"phone": "901000001", "pin": "1234"}).status_code == 429


def test_checkin_needs_login(client):
    assert client.post("/api/checkin", json={"zone": 1, "code": "1111"}).status_code == 401


def test_mod_required(client):
    assert client.get("/api/mod/stats").status_code == 401
    assert client.post("/api/mod/login", json={"login": "aziz", "password": "bad"}).status_code == 401


def test_mod_reset_pin(client, mod):
    c = register(client, 1, stamps=2)
    uid = mod.get("/api/mod/users?q=Ism1").json()["users"][0]["id"]
    assert mod.post(f"/api/mod/users/{uid}/reset-pin", json={"pin": "7777"}).status_code == 200
    assert c.get("/api/me").status_code == 401  # eski sessiya tugagan
    from app.main import app
    c2 = TestClient(app)
    assert c2.post("/api/login", json={"phone": "901000001", "pin": "1234"}).status_code == 401
    assert c2.post("/api/login", json={"phone": "901000001", "pin": "7777"}).status_code == 200
    assert c2.get("/api/me").json()["stamps"] == [1, 2]


def test_mod_manual_stamp_block_stats(client, mod):
    register(client, 1, stamps=0)
    uid = mod.get("/api/mod/users").json()["users"][0]["id"]
    mod.post(f"/api/mod/users/{uid}/stamp", json={"zone": 4, "on": True})
    s = mod.get("/api/mod/stats").json()
    assert s["total_users"] == 1 and s["by_zone"]["4"] == 1
    mod.post(f"/api/mod/users/{uid}/block", json={"blocked": True})
    from app.main import app
    c = TestClient(app)
    assert c.post("/api/login", json={"phone": "901000001", "pin": "1234"}).status_code == 403
    assert any(a["action"] == "block" for a in mod.get("/api/mod/audit").json()["log"])


def test_exports(client, mod):
    register(client, 1, stamps=9)
    assert "ishtirokchilar" in mod.get("/api/mod/export/users.csv").headers["content-disposition"]
    assert mod.get("/api/mod/export/rounds.csv").status_code == 200


def test_login_ip_limit(client, monkeypatch):
    from app import config
    monkeypatch.setattr(config, "LOGIN_IP_MAX_FAILS", 3)
    for i in range(3):
        register(client, i + 1, stamps=0)
    from app.main import app
    c = TestClient(app)
    for i in range(3):   # har xil telefonlarga PIN "purkash"
        assert c.post("/api/login", json={"phone": f"90100{i + 1:04d}", "pin": "0000"}).status_code == 401
    assert c.post("/api/login", json={"phone": "901000001", "pin": "1234"}).status_code == 429


def test_csv_formula_injection_escaped(client, mod):
    from app.main import app
    c = TestClient(app)
    assert c.post("/api/register", json={"first_name": "=HYPERLINK(1)", "last_name": "@x",
                                         "phone": "901112233", "pin": "1234", "visitor_type_id": 1}).status_code == 200
    body = mod.get("/api/mod/export/users.csv").text
    assert "'=HYPERLINK(1)" in body and "'@x" in body


def test_config_check_rejects_missing_codes(monkeypatch):
    import pytest
    from app import config
    from app.main import check_config
    monkeypatch.setattr(config, "ZONE_CODES", {1: "1111"})
    with pytest.raises(RuntimeError):
        check_config()


def test_xlsx_export_styled(client, mod):
    import io
    from openpyxl import load_workbook
    from app.main import app
    c = TestClient(app)
    assert c.post("/api/register", json={"first_name": "=1+1", "last_name": "B", "phone": "901112233",
                                         "pin": "1234", "visitor_type_id": 3}).status_code == 200
    r = mod.get("/api/mod/export/users.xlsx")
    assert r.status_code == 200 and "ishtirokchilar.xlsx" in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A1"].value == "ID" and ws["A1"].font.bold and ws.freeze_panes == "A2" and ws.auto_filter.ref
    assert ws["C2"].value == "=1+1" and ws["C2"].data_type == "s"   # formula emas, matn
    assert ws["F2"].value == "Талаба"
    assert mod.get("/api/mod/export/rounds.xlsx").status_code == 200
